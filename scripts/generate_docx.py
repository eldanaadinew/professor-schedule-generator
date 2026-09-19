#!/usr/bin/env python3
import csv
import sys
from collections import defaultdict
from datetime import datetime
from pathlib import Path

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn
from docx.shared import Inches, Pt


def display_time(value):
    return datetime.strptime(value, "%H:%M").strftime("%-I:%M %p")


def shade(cell, fill):
    cell._tc.get_or_add_tcPr().append(parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill}"/>'))


def format_cell(cell, align=WD_ALIGN_PARAGRAPH.LEFT, header=False):
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    properties = cell._tc.get_or_add_tcPr()
    margins = properties.first_child_found_in("w:tcMar")
    if margins is None:
        margins = OxmlElement("w:tcMar")
        properties.append(margins)
    for edge in ("top", "left", "bottom", "right"):
        node = margins.find(qn(f"w:{edge}"))
        if node is None:
            node = OxmlElement(f"w:{edge}")
            margins.append(node)
        node.set(qn("w:w"), "100")
        node.set(qn("w:type"), "dxa")
    for paragraph in cell.paragraphs:
        paragraph.alignment = align
        paragraph.paragraph_format.space_after = Pt(0)
        for run in paragraph.runs:
            run.font.name = "Arial"
            run.font.size = Pt(10.5)
            if header:
                run.bold = True
                run.font.color.rgb = __import__("docx").shared.RGBColor(255, 255, 255)


def set_table_borders(table):
    properties = table._tbl.tblPr
    borders = properties.first_child_found_in("w:tblBorders")
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        properties.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        node = borders.find(qn(f"w:{edge}"))
        if node is None:
            node = OxmlElement(f"w:{edge}")
            borders.append(node)
        node.set(qn("w:val"), "single")
        node.set(qn("w:sz"), "6")
        node.set(qn("w:color"), "D9D9D9")


def main():
    if len(sys.argv) != 4:
        raise SystemExit("Usage: generate_docx.py class.csv office.csv output.docx")
    with open(sys.argv[1], newline="", encoding="utf-8") as stream:
        classes = list(csv.DictReader(stream))
    with open(sys.argv[2], newline="", encoding="utf-8") as stream:
        office = list(csv.DictReader(stream))

    faculty = defaultdict(lambda: {"classes": [], "office": []})
    for row in classes:
        faculty[row["Faculty"]]["classes"].append(row)
    for row in office:
        faculty[row["Faculty"]]["office"].append(row)

    doc = Document()
    section = doc.sections[0]
    section.top_margin = section.bottom_margin = Inches(0.65)
    section.left_margin = section.right_margin = Inches(0.65)
    styles = doc.styles
    styles["Normal"].font.name = "Arial"
    styles["Normal"].font.size = Pt(11)
    styles["Title"].font.name = "Arial"
    styles["Title"].font.color.rgb = __import__("docx").shared.RGBColor(0, 0, 0)
    style_border = styles["Title"].element.get_or_add_pPr().find(qn("w:pBdr"))
    if style_border is not None:
        styles["Title"].element.get_or_add_pPr().remove(style_border)

    for index, (name, data) in enumerate(faculty.items()):
        if index:
            doc.add_page_break()
        title = doc.add_paragraph(style="Title")
        title.alignment = WD_ALIGN_PARAGRAPH.CENTER
        title_run = title.add_run("Faculty Teaching and Office Hours Schedule")
        title_run.font.size = Pt(22)
        title_properties = title._p.get_or_add_pPr()
        title_border = title_properties.find(qn("w:pBdr"))
        if title_border is not None:
            title_properties.remove(title_border)
        sample = data["classes"][0] if data["classes"] else data["office"][0]
        subtitle = doc.add_paragraph(f'{name}\n{sample["Department"]} | Office {data["office"][0]["Office"]}')
        subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
        subtitle.runs[0].bold = True

        heading = doc.add_paragraph()
        heading.add_run("CLASS SCHEDULE").bold = True
        table = doc.add_table(rows=1, cols=3)
        table.style = "Table Grid"
        table.autofit = False
        set_table_borders(table)
        headers = ["Course", "Times", "Location"]
        for i, label in enumerate(headers):
            table.rows[0].cells[i].text = label
            shade(table.rows[0].cells[i], "1F4E78")
            format_cell(table.rows[0].cells[i], WD_ALIGN_PARAGRAPH.CENTER, header=True)
        for row in data["classes"]:
            cells = table.add_row().cells
            cells[0].text = f'{row["Course"]}\n({row["Title"]})'
            cells[1].text = f'{row["Days"]} {display_time(row["Start"])} to {display_time(row["End"])}'
            cells[2].text = row["Room"]
            format_cell(cells[0])
            format_cell(cells[1])
            format_cell(cells[2], WD_ALIGN_PARAGRAPH.CENTER)
        table.columns[0].width = Inches(2.15)
        table.columns[1].width = Inches(4.15)
        table.columns[2].width = Inches(0.9)

        doc.add_paragraph()
        heading = doc.add_paragraph()
        heading.add_run("OFFICE HOURS").bold = True
        box = doc.add_table(rows=0, cols=1)
        box.style = "Table Grid"
        box.autofit = False
        set_table_borders(box)
        for row in data["office"]:
            cell = box.add_row().cells[0]
            cell.text = (f'{row["Day"]}: {display_time(row["Start"])} - '
                         f'{display_time(row["End"])} ({row["Related Course"]})')
            cell.paragraphs[0].runs[0].bold = True
            cell.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
            format_cell(cell, WD_ALIGN_PARAGRAPH.CENTER)
        note = doc.add_paragraph("NOTE: Instructors are not to be interrupted during class times.")
        note.runs[0].bold = True
        note.runs[0].font.size = Pt(9)

    Path(sys.argv[3]).parent.mkdir(parents=True, exist_ok=True)
    doc.save(sys.argv[3])


if __name__ == "__main__":
    main()
